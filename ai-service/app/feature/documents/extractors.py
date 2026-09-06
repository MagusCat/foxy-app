import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath

from app.platform.errors import UnsupportedType


@dataclass(frozen=True)
class Limits:
    max_decompressed_bytes: int
    max_chars: int
    max_pdf_pages: int


IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


def suffix_of(storage_path: str) -> str:
    return PurePosixPath(storage_path).suffix.lower()


def is_image(suffix: str) -> bool:
    return suffix in IMAGE_MIME


def extract(blob: bytes, suffix: str, limits: Limits) -> str:
    match suffix:
        case ".pdf":
            return _pdf(blob, limits)
        case ".docx":
            return _docx(blob, limits)
        case ".pptx":
            return _pptx(blob, limits)
        case ".txt" | ".md" | ".markdown" | "":
            return blob[: limits.max_chars].decode("utf-8", errors="replace").strip()
        case _:
            raise UnsupportedType(f"no hay extractor para archivos {suffix or 'sin extensión'}")


def _guard_zip(blob: bytes, limits: Limits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            declared = sum(entry.file_size for entry in z.infolist())
    except zipfile.BadZipFile as e:
        raise UnsupportedType("el archivo no es un ZIP válido") from e
    if declared > limits.max_decompressed_bytes:
        raise UnsupportedType(
            f"el archivo se expande a {declared // 1024 // 1024} MB, por encima del "
            f"máximo de {limits.max_decompressed_bytes // 1024 // 1024} MB"
        )


def _pdf(blob: bytes, limits: Limits) -> str:
    from pypdf import PdfReader

    pages = PdfReader(io.BytesIO(blob)).pages
    parts: list[str] = []
    total = 0
    for page in list(pages)[: limits.max_pdf_pages]:
        text = page.extract_text() or ""
        parts.append(text)
        total += len(text)
        if total >= limits.max_chars:
            break
    return "\n\n".join(parts).strip()


def _docx(blob: bytes, limits: Limits) -> str:
    import docx

    _guard_zip(blob, limits)
    return "\n".join(p.text for p in docx.Document(io.BytesIO(blob)).paragraphs).strip()


def _pptx(blob: bytes, limits: Limits) -> str:
    import pptx

    _guard_zip(blob, limits)
    slides = pptx.Presentation(io.BytesIO(blob)).slides
    return "\n\n".join(
        "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        for slide in slides
    ).strip()
